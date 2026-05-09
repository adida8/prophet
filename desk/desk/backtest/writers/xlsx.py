"""Workbook writer.

Open `desk_backtest.xlsx`, clear the Snapshots and Match Universe rows
below the header, write fresh rows from the SnapshotRow array, preserve
the formula columns (Brier and verdict_resolved auto-recompute when
Excel opens the file).
"""

from __future__ import annotations

import logging
from pathlib import Path

from openpyxl import load_workbook

from desk.backtest.historical.markets import HistoricalMatch
from desk.backtest.replay import SnapshotRow

log = logging.getLogger("desk.backtest.writers.xlsx")


# Column indexes in the Snapshots tab (1-based for openpyxl). Mirrors the
# scaffold header row exactly.
COL_SNAPSHOTS = {
    "match_id":             1,
    "window":               2,
    "elo_a":                3,
    "elo_b":                4,
    "host_bonus":           5,
    "altitude_m":           6,
    "weather_factor":       7,
    "injury_factor":        8,
    "p_a":                  9,
    "p_draw":              10,
    "p_b":                 11,
    "market_p_a":          12,
    "market_p_draw":       13,
    "market_p_b":          14,
    "actual_a":            15,
    "actual_draw":         16,
    "actual_b":            17,
    "verdict_state":       18,
    "verdict_side":        19,
    "verdict_market_venue":20,
    "verdict_edge_pp":     21,
    "brier_score":         22,    # formula
    "verdict_resolved":    23,    # formula
    "is_example":          24,
}

COL_UNIVERSE = {
    "match_id":      1,
    "tournament":    2,
    "stage":         3,
    "kickoff_utc":   4,
    "team_a":        5,
    "team_b":        6,
    "venue_city":    7,
    "venue_country": 8,
    "result_a":      9,
    "result_b":     10,
    "winner":       11,
    "is_example":   12,
}


def _clear_below_header(ws) -> None:
    """Wipe every cell below row 1.

    `ws.delete_rows(2, n)` plus openpyxl's lazy `max_row` calculation can
    leave stragglers — explicitly null out values in the data band so
    callers see a clean tab. Merged cells are skipped (they're read-only
    via the public Cell API and don't carry stale data after delete_rows
    anyway).
    """
    from openpyxl.cell.cell import MergedCell

    last = ws.max_row
    if last <= 1:
        return
    for row in ws.iter_rows(min_row=2, max_row=last):
        for cell in row:
            if isinstance(cell, MergedCell):
                continue
            cell.value = None
    ws.delete_rows(2, last)


def _brier_formula(row: int) -> str:
    return f"=(I{row}-O{row})^2+(J{row}-P{row})^2+(K{row}-Q{row})^2"


def _verdict_resolved_formula(row: int) -> str:
    return (
        f'=IF(B{row}<>"KO","",'
        f'IF(R{row}="pick",'
        f'IF(VLOOKUP(A{row},\'Match Universe\'!A:K,11,FALSE)=S{row},"hit","miss"),'
        f'IF(R{row}="pass","n/a","n/a")))'
    )


def write_xlsx(
    *,
    workbook_path: Path,
    matches: list[HistoricalMatch],
    snapshots: list[SnapshotRow],
) -> None:
    """Open the workbook, clear & rewrite Snapshots and Match Universe."""
    wb = load_workbook(workbook_path)

    # ── Match Universe ───────────────────────────────────────────────
    mu = wb["Match Universe"]
    _clear_below_header(mu)
    for i, m in enumerate(matches, start=2):
        mu.cell(row=i, column=COL_UNIVERSE["match_id"],      value=m.match_id)
        mu.cell(row=i, column=COL_UNIVERSE["tournament"],    value=m.tournament)
        mu.cell(row=i, column=COL_UNIVERSE["stage"],         value=m.stage)
        mu.cell(row=i, column=COL_UNIVERSE["kickoff_utc"],
                value=m.kickoff_utc.strftime("%Y-%m-%d %H:%M"))
        mu.cell(row=i, column=COL_UNIVERSE["team_a"],        value=m.team_a)
        mu.cell(row=i, column=COL_UNIVERSE["team_b"],        value=m.team_b)
        mu.cell(row=i, column=COL_UNIVERSE["venue_city"],    value=m.venue_city)
        mu.cell(row=i, column=COL_UNIVERSE["venue_country"], value=m.venue_country)
        mu.cell(row=i, column=COL_UNIVERSE["result_a"],      value=m.goals_a)
        mu.cell(row=i, column=COL_UNIVERSE["result_b"],      value=m.goals_b)
        mu.cell(row=i, column=COL_UNIVERSE["winner"],        value=m.winner_90min)
        mu.cell(row=i, column=COL_UNIVERSE["is_example"],
                value=f"BACKTEST {m.competition} {m.season}")

    # ── Snapshots ────────────────────────────────────────────────────
    sn = wb["Snapshots"]
    _clear_below_header(sn)
    for i, s in enumerate(snapshots, start=2):
        sn.cell(row=i, column=COL_SNAPSHOTS["match_id"],             value=s.match_id)
        sn.cell(row=i, column=COL_SNAPSHOTS["window"],               value=s.window)
        sn.cell(row=i, column=COL_SNAPSHOTS["elo_a"],                value=round(s.elo_a, 1))
        sn.cell(row=i, column=COL_SNAPSHOTS["elo_b"],                value=round(s.elo_b, 1))
        sn.cell(row=i, column=COL_SNAPSHOTS["host_bonus"],           value=round(s.host_bonus_pp, 1))
        sn.cell(row=i, column=COL_SNAPSHOTS["altitude_m"],           value=s.altitude_m)
        sn.cell(row=i, column=COL_SNAPSHOTS["weather_factor"],       value=s.weather_factor)
        sn.cell(row=i, column=COL_SNAPSHOTS["injury_factor"],        value=s.injury_factor)
        sn.cell(row=i, column=COL_SNAPSHOTS["p_a"],                  value=s.p_a)
        sn.cell(row=i, column=COL_SNAPSHOTS["p_draw"],               value=s.p_draw)
        sn.cell(row=i, column=COL_SNAPSHOTS["p_b"],                  value=s.p_b)
        sn.cell(row=i, column=COL_SNAPSHOTS["market_p_a"],           value=s.market_p_a)
        sn.cell(row=i, column=COL_SNAPSHOTS["market_p_draw"],        value=s.market_p_draw)
        sn.cell(row=i, column=COL_SNAPSHOTS["market_p_b"],           value=s.market_p_b)
        sn.cell(row=i, column=COL_SNAPSHOTS["actual_a"],             value=s.actual_a)
        sn.cell(row=i, column=COL_SNAPSHOTS["actual_draw"],          value=s.actual_draw)
        sn.cell(row=i, column=COL_SNAPSHOTS["actual_b"],             value=s.actual_b)
        sn.cell(row=i, column=COL_SNAPSHOTS["verdict_state"],        value=s.verdict_state)
        sn.cell(row=i, column=COL_SNAPSHOTS["verdict_side"],         value=s.verdict_side or "")
        sn.cell(row=i, column=COL_SNAPSHOTS["verdict_market_venue"], value=s.verdict_market_venue or "")
        sn.cell(row=i, column=COL_SNAPSHOTS["verdict_edge_pp"],      value=s.verdict_edge_pp)
        sn.cell(row=i, column=COL_SNAPSHOTS["brier_score"],          value=_brier_formula(i))
        sn.cell(row=i, column=COL_SNAPSHOTS["verdict_resolved"],     value=_verdict_resolved_formula(i))
        sn.cell(row=i, column=COL_SNAPSHOTS["is_example"],           value=s.is_example)

    wb.save(workbook_path)
    log.info("wrote %d snapshots + %d matches to %s",
             len(snapshots), len(matches), workbook_path)
