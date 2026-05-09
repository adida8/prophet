"""Writer tests — workbook round-trip + dashboard marker behaviour."""

from __future__ import annotations

import shutil
from datetime import date, datetime, timezone
from pathlib import Path

import pytest
from openpyxl import load_workbook

from desk.backtest.historical.markets import HistoricalMatch
from desk.backtest.replay import SnapshotRow
from desk.backtest.writers.dashboard import write_dashboard
from desk.backtest.writers.xlsx import write_xlsx

PROJECT_ROOT = Path(__file__).resolve().parents[3]
SCAFFOLD_XLSX = PROJECT_ROOT / "desk_backtest.xlsx"
SCAFFOLD_HTML = PROJECT_ROOT / "desk_backtest_dashboard.html"


def _sample_match() -> HistoricalMatch:
    return HistoricalMatch(
        match_id="wc22-c-arg-sau-20221122",
        tournament="FIFA World Cup 2022",
        competition="wc",
        season="2022",
        stage="Group C",
        kickoff_utc=datetime(2022, 11, 22, 10, 0, tzinfo=timezone.utc),
        team_a="Argentina", team_a_iso3="arg",
        team_b="Saudi Arabia", team_b_iso3="ksa",
        venue_city="Lusail", venue_country="QA", venue_altitude_m=8.0,
        goals_a=1, goals_b=2, winner_90min="b",
        close_a=1.18, close_draw=8.0, close_b=15.0, close_source="Pinnacle",
    )


def _sample_row(window: str = "KO") -> SnapshotRow:
    return SnapshotRow(
        match_id="wc22-c-arg-sau-20221122",
        window=window,                      # type: ignore[arg-type]
        asof=datetime(2022, 11, 22, 9, 0, tzinfo=timezone.utc),
        elo_a=2142, elo_b=1635, host_bonus_pp=0, altitude_m=8,
        weather_factor=1, injury_factor=1,
        p_a=0.82, p_draw=0.12, p_b=0.06,
        market_p_a=0.85, market_p_draw=0.10, market_p_b=0.05,
        actual_a=0, actual_draw=0, actual_b=1 if window == "KO" else 0,
        verdict_state="pass", verdict_side=None,
        verdict_market_venue=None, verdict_edge_pp=-3.0,
        competition="wc", season="2022",
        tournament="FIFA World Cup 2022", stage="Group C",
        kickoff_utc=datetime(2022, 11, 22, 10, 0, tzinfo=timezone.utc),
        team_a="Argentina", team_b="Saudi Arabia",
        is_example="BACKTEST wc 2022",
    )


# ── Workbook ─────────────────────────────────────────────────────────

def test_xlsx_round_trip_preserves_formula_columns(tmp_path: Path) -> None:
    target = tmp_path / "wb.xlsx"
    shutil.copy(SCAFFOLD_XLSX, target)

    write_xlsx(
        workbook_path=target,
        matches=[_sample_match()],
        snapshots=[_sample_row("T-5"), _sample_row("KO")],
    )

    wb = load_workbook(target, data_only=False)
    sn = wb["Snapshots"]
    # We can't assert sn.max_row strictly — openpyxl tracks max_row via
    # styled cells too, and the scaffold ships with formatting on rows
    # 2..26 that survives clears. Assert on actual data instead: rows 2
    # and 3 carry the new snapshots; rows 4+ are empty.
    assert sn.cell(row=2, column=1).value == "wc22-c-arg-sau-20221122"
    assert sn.cell(row=3, column=1).value == "wc22-c-arg-sau-20221122"
    assert sn.cell(row=4, column=1).value is None

    # Brier formula must be a real formula, not a number we computed.
    brier_cell = sn.cell(row=2, column=22)
    assert isinstance(brier_cell.value, str) and brier_cell.value.startswith("=")
    # KO row must reference VLOOKUP into Match Universe in verdict_resolved.
    resolved_cell = sn.cell(row=3, column=23)
    assert "VLOOKUP" in resolved_cell.value


def test_xlsx_replaces_example_rows(tmp_path: Path) -> None:
    target = tmp_path / "wb.xlsx"
    shutil.copy(SCAFFOLD_XLSX, target)

    write_xlsx(
        workbook_path=target,
        matches=[_sample_match()],
        snapshots=[_sample_row("KO")],
    )
    wb = load_workbook(target)
    # is_example column = 24 in Snapshots tab
    assert wb["Snapshots"].cell(row=2, column=24).value == "BACKTEST wc 2022"
    # is_example column = 12 in Match Universe
    assert wb["Match Universe"].cell(row=2, column=12).value == "BACKTEST wc 2022"


# ── Dashboard ────────────────────────────────────────────────────────

def test_dashboard_inserts_markers_and_replaces_kpis(tmp_path: Path) -> None:
    target = tmp_path / "dash.html"
    shutil.copy(SCAFFOLD_HTML, target)

    write_dashboard(
        dashboard_path=target,
        snapshots=[_sample_row("KO")],
        competitions=["FIFA World Cup 2022"],
        run_date=date(2026, 5, 9),
    )

    body = target.read_text(encoding="utf-8")
    assert "BACKTEST_DATA_START:masthead"   in body
    assert "BACKTEST_DATA_START:disclaimer" in body
    assert "BACKTEST_DATA_START:kpis"       in body
    assert "BACKTEST_DATA_START:match_table" in body
    assert "BACKTEST_DATA_START:bins"       in body

    # KPI strip carries computed values.
    assert "Matches tested" in body
    # No more "Scaffold — illustrative" badge.
    assert "Scaffold — illustrative" not in body
    # Disclaimer carries the run-date.
    assert "2026-05-09" in body


def test_dashboard_idempotent_second_call_is_safe(tmp_path: Path) -> None:
    target = tmp_path / "dash.html"
    shutil.copy(SCAFFOLD_HTML, target)

    write_dashboard(
        dashboard_path=target,
        snapshots=[_sample_row("KO")],
        competitions=["FIFA World Cup 2022"],
        run_date=date(2026, 5, 9),
    )
    after_first = target.read_text(encoding="utf-8")

    write_dashboard(
        dashboard_path=target,
        snapshots=[_sample_row("KO")],
        competitions=["FIFA World Cup 2022"],
        run_date=date(2026, 5, 9),
    )
    after_second = target.read_text(encoding="utf-8")

    # Same input → same output (idempotent).
    assert after_first == after_second
    # No marker duplication.
    assert after_second.count("BACKTEST_DATA_START:kpis") == 1
